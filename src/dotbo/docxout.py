"""Write the consolidated Word document in the reference format.

Rebuilt from ``Blocking orders_May.docx``:

* Title - Arial 12pt, bold, small caps, centred, single bottom border.
* One blank Arial 10pt paragraph.
* Per order - a plain Arial 10pt justified paragraph carrying the same text as
  the PDF's filename, then the URLs as a numbered list that **restarts at 1**
  (the reference carries one ``numId`` per order), then a blank paragraph.
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .naming import Order, compilation_name

FONT = "Arial"
BODY_PT = 10
TITLE_PT = 12
HYPERLINK_BLUE = RGBColor(0x05, 0x63, 0xC1)


def build(orders: list[Order], year: int, month: int, path: str | Path) -> Path:
    """Write the compilation and return the path."""
    path = Path(path)
    document = docx.Document()
    _page_setup(document)
    _title(document, compilation_name(year, month))
    _blank(document)

    for order in orders:
        _heading(document, order.stem)
        num_id = _new_numbering(document)
        for url in order.urls:
            _url_item(document, url, num_id)
        for note in _notes(order):
            _note(document, note)
        _blank(document)

    document.save(path)
    return path


# ---------------------------------------------------------------- pieces


def _page_setup(document) -> None:
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)
    style.paragraph_format.space_after = Pt(0)


def _title(document, text: str) -> None:
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    _font(run, TITLE_PT, bold=True, small_caps=True)
    _bottom_border(para)


def _heading(document, text: str) -> None:
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.space_after = Pt(0)
    _font(para.add_run(text), BODY_PT)


def _blank(document) -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    _font(para.add_run(""), BODY_PT)


def _note(document, text: str) -> None:
    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    _font(run, BODY_PT - 1)
    run.italic = True


def _url_item(document, url: str, num_id: int) -> None:
    para = document.add_paragraph(style="List Paragraph")
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.space_after = Pt(0)
    _apply_numbering(para, num_id)
    _hyperlink(para, url)


def _font(run, size: int, *, bold: bool = False, small_caps: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if small_caps:
        el = rpr.makeelement(qn("w:smallCaps"), {})
        rpr.append(el)


def _bottom_border(para) -> None:
    ppr = para._p.get_or_add_pPr()
    borders = ppr.makeelement(qn("w:pBdr"), {})
    bottom = borders.makeelement(
        qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "1", qn("w:color"): "auto"},
    )
    borders.append(bottom)
    ppr.append(borders)


def _apply_numbering(para, num_id: int) -> None:
    ppr = para._p.get_or_add_pPr()
    numpr = ppr.makeelement(qn("w:numPr"), {})
    ilvl = numpr.makeelement(qn("w:ilvl"), {qn("w:val"): "0"})
    nid = numpr.makeelement(qn("w:numId"), {qn("w:val"): str(num_id)})
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)
    ind = ppr.makeelement(qn("w:ind"), {qn("w:left"): "720", qn("w:hanging"): "360"})
    ppr.append(ind)


def _hyperlink(para, url: str) -> None:
    """A real Word hyperlink, styled like the reference document."""
    target = url.strip()
    href = target if target.lower().startswith(("http://", "https://")) else "http://" + target
    try:
        r_id = para.part.relate_to(
            href,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
    except Exception:  # noqa: BLE001 - an unusable URL still belongs in the list
        _font(para.add_run(target), BODY_PT)
        return

    link = para._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    run = para._p.makeelement(qn("w:r"), {})
    rpr = run.makeelement(qn("w:rPr"), {})
    rfonts = rpr.makeelement(
        qn("w:rFonts"), {qn("w:ascii"): FONT, qn("w:hAnsi"): FONT, qn("w:cs"): FONT}
    )
    color = rpr.makeelement(qn("w:color"), {qn("w:val"): "0563C1"})
    underline = rpr.makeelement(qn("w:u"), {qn("w:val"): "single"})
    size = rpr.makeelement(qn("w:sz"), {qn("w:val"): str(BODY_PT * 2)})
    size_cs = rpr.makeelement(qn("w:szCs"), {qn("w:val"): str(BODY_PT * 2)})
    for el in (rfonts, color, underline, size, size_cs):
        rpr.append(el)
    text = run.makeelement(qn("w:t"), {})
    text.text = target
    run.append(rpr)
    run.append(text)
    link.append(run)
    para._p.append(link)


# ---------------------------------------------------------------- numbering


def _numbering(document):
    part = document.part.numbering_part
    return part.element


def _new_numbering(document) -> int:
    """Create a fresh decimal list so each order's numbering restarts at 1."""
    numbering = _numbering(document)
    abstract_ids = [
        int(e.get(qn("w:abstractNumId")))
        for e in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = numbering.makeelement(
        qn("w:abstractNum"), {qn("w:abstractNumId"): str(abstract_id)}
    )
    multi = abstract.makeelement(qn("w:multiLevelType"), {qn("w:val"): "hybridMultilevel"})
    abstract.append(multi)
    for level in range(9):
        lvl = abstract.makeelement(qn("w:lvl"), {qn("w:ilvl"): str(level)})
        start = lvl.makeelement(qn("w:start"), {qn("w:val"): "1"})
        fmt = lvl.makeelement(
            qn("w:numFmt"),
            {qn("w:val"): "decimal" if level % 3 == 0 else ("lowerLetter" if level % 3 == 1 else "lowerRoman")},
        )
        text = lvl.makeelement(qn("w:lvlText"), {qn("w:val"): "%%%d." % (level + 1)})
        jc = lvl.makeelement(qn("w:lvlJc"), {qn("w:val"): "left"})
        ppr = lvl.makeelement(qn("w:pPr"), {})
        ind = ppr.makeelement(
            qn("w:ind"),
            {qn("w:left"): str(720 * (level + 1)), qn("w:hanging"): "360"},
        )
        ppr.append(ind)
        for el in (start, fmt, text, jc, ppr):
            lvl.append(el)
        abstract.append(lvl)

    num = numbering.makeelement(qn("w:num"), {qn("w:numId"): str(num_id)})
    ref = num.makeelement(qn("w:abstractNumId"), {qn("w:val"): str(abstract_id)})
    num.append(ref)

    # abstractNum elements must all precede num elements.
    last_abstract = numbering.findall(qn("w:abstractNum"))
    if last_abstract:
        last_abstract[-1].addnext(abstract)
    else:
        numbering.insert(0, abstract)
    numbering.append(num)
    return num_id


def _notes(order: Order) -> list[str]:
    """Only what changes the legal meaning of the entry.

    How the text was read - OCR, which page span, whether the count matched to
    the digit - belongs in the QA report, not in the compilation.  What belongs
    here is anything a reader of the document would be misled by without it.
    """
    ex = order.extraction
    notes: list[str] = []
    if order.error:
        notes.append(f"[Not processed: {order.error}]")
        return notes
    if ex is None:
        return notes
    if ex.source == "court_order":
        notes.append(
            "[URLs above taken from the court order/judgement - no list was annexed "
            "to the DoT instruction.]"
        )
    if order.letter is not None and order.letter.action == "unblock":
        notes.append("[This is an UNBLOCKING instruction.]")
    if not ex.urls:
        # The May reference does exactly this for its order 19.
        notes.append(
            "[No website list could be extracted automatically. Source PDF: "
            f"{order.attachment.pdf_url}]"
        )
    return notes
